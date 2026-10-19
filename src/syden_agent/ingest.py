"""Read a folder of documents. Fail independently per file; never invent extracted data."""
from __future__ import annotations

import datetime as dt
import json
import os
import re
import sqlite3
from pathlib import Path
from typing import Any

from .db import mark_document, replace_document

SUPPORTED = {".pdf", ".xlsx", ".xlsm", ".xls", ".docx"}
KNOWN_UNSUPPORTED = {".xlsb"}


def _discover(directory: Path) -> list[Path]:
    """Walk locally without following linked directories or linked files."""
    found: list[Path] = []
    for root, dirs, files in os.walk(directory, followlinks=False):
        dirs[:] = sorted(name for name in dirs if not (Path(root) / name).is_symlink())
        for name in sorted(files):
            path = Path(root) / name
            if name.startswith("~$") or path.is_symlink():
                continue
            if path.resolve().is_relative_to(directory) and path.is_file():
                found.append(path)
    return sorted(found)


def _cell_value(value: Any) -> Any:
    if isinstance(value, (dt.date, dt.datetime, dt.time)):
        return value.isoformat()
    if isinstance(value, (int, float, str, bool)) or value is None:
        return value
    return str(value)


def _chunks(text: str, size: int = 1500) -> list[str]:
    """Page-aware paragraph chunks; overlap is added in the Codex hardening phase."""
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    output: list[str] = []
    buffer = ""
    for paragraph in paragraphs:
        if len(buffer) + len(paragraph) <= size:
            buffer = (buffer + "\n\n" + paragraph).strip()
        else:
            if buffer:
                output.append(buffer)
            while len(paragraph) > size:
                output.append(paragraph[:size])
                paragraph = paragraph[size:]
            buffer = paragraph
    if buffer:
        output.append(buffer)
    return output


def _ingest_pdf(conn: sqlite3.Connection, path: Path, doc_id: int) -> tuple[int, list[str]]:
    import pdfplumber
    from openpyxl.utils import get_column_letter

    total = 0
    empty_pages: list[int] = []
    with pdfplumber.open(path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if not text.strip():
                empty_pages.append(page_number)
            for chunk_index, chunk in enumerate(_chunks(text)):
                conn.execute(
                    "INSERT INTO passages(document_id,page,chunk,text) VALUES(?,?,?,?)",
                    (doc_id, page_number, chunk_index, chunk),
                )
                total += 1
            tables = page.extract_tables() or []
            for table_number, table in enumerate(tables, 1):
                for row_number, values in enumerate(table, 1):
                    cells = {get_column_letter(index + 1): {"value": value.strip()}
                             for index, value in enumerate(values)
                             if isinstance(value, str) and value.strip()}
                    if cells:
                        conn.execute(
                            "INSERT INTO pdf_boq_rows(document_id,page,table_num,row_num,"
                            "cells_json,method) VALUES(?,?,?,?,?,?)",
                            (doc_id, page_number, table_number, row_number,
                             json.dumps(cells, ensure_ascii=False), "detected_table"),
                        )
            lines = [line.strip() for line in text.splitlines() if line.strip()]
            looks_like_boq = bool(re.search(
                r"\b(?:BOQ|bill of quantities)\b",
                text[:1000] + " " + path.stem.replace("_", " "), re.IGNORECASE))
            has_header = any(re.search(r"\b(?:item|description|equipment)\b.*"
                                       r"\b(?:amount|total|price|cost)\b", line,
                                       re.IGNORECASE) for line in lines[:30])
            if not tables and text and (looks_like_boq or has_header):
                # Conservative fallback for selectable-text BOQs lacking PDF table borders.
                parsed = []
                for line_number, line in enumerate(lines, 1):
                    match = re.fullmatch(
                        r"(.+?)\s+(?:(SAR|AED|USD|EUR|GBP)\s+)?"
                        r"(\(?[-+]?\d[\d,]*(?:\.\d+)?\)?)"
                        r"(?:\s+(SAR|AED|USD|EUR|GBP))?", line, re.IGNORECASE)
                    if match and not re.search(r"\b(?:date|year|revision)\b", match.group(1),
                                               re.IGNORECASE):
                        parsed.append((line_number, match))
                if parsed:
                    declared = re.search(r"\bcurrency\s*[:=-]\s*(SAR|AED|USD|EUR|GBP)\b",
                                         text, re.IGNORECASE)
                    header = "Amount" + (f" ({declared.group(1).upper()})" if declared else "")
                    conn.execute(
                        "INSERT INTO pdf_boq_rows(document_id,page,table_num,row_num,cells_json,method) "
                        "VALUES(?,?,?,?,?,?)",
                        (doc_id, page_number, 1, 0,
                         json.dumps({"A": {"value": "Item"}, "B": {"value": header}}),
                         "text_lines"),
                    )
                    for line_number, match in parsed:
                        amount = match.group(3)
                        code = match.group(2) or match.group(4)
                        if code:
                            amount = f"{code.upper()} {amount}"
                        cells = {"A": {"value": match.group(1)}, "B": {"value": amount}}
                        conn.execute(
                            "INSERT INTO pdf_boq_rows(document_id,page,table_num,row_num,"
                            "cells_json,method) VALUES(?,?,?,?,?,?)",
                            (doc_id, page_number, 1, line_number,
                             json.dumps(cells, ensure_ascii=False), "text_lines"),
                        )
        conn.commit()
    if total == 0:
        raise ValueError("PDF has no extractable text; it may be scanned or image-only")
    warnings = ([f"Pages without extractable text (possibly scanned): {empty_pages[:20]}"
                 + (" and more" if len(empty_pages) > 20 else "")]
                if empty_pages else [])
    return total, warnings


def _ingest_workbook(conn: sqlite3.Connection, path: Path, doc_id: int) -> int:
    from openpyxl import load_workbook
    # Never execute macros. Read-only mode minimizes memory pressure on large workbooks.
    # data_only reads Excel's cached result; it does NOT calculate formulas.
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    formula_wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    total = 0
    try:
        for sheet in wb.worksheets:
            formula_sheet = formula_wb[sheet.title]
            cached_rows = sheet.iter_rows()
            raw_rows = formula_sheet.iter_rows()
            for row_number, (cached_row, raw_row) in enumerate(zip(cached_rows, raw_rows), 1):
                cells = {}
                for cached, raw in zip(cached_row, raw_row):
                    value = _cell_value(cached.value)
                    if value is None and raw.data_type != "f":
                        continue
                    entry: dict[str, Any] = {"value": value, "number_format": raw.number_format}
                    if raw.data_type == "f":
                        entry["formula"] = str(raw.value)
                        if value is None:
                            entry["warning"] = "Formula has no cached result"
                    cells[raw.column_letter] = entry
                if cells:
                    conn.execute(
                        "INSERT INTO sheet_rows(document_id,sheet,row_num,cells_json) VALUES(?,?,?,?)",
                        (doc_id, sheet.title, row_number, json.dumps(cells, ensure_ascii=False)),
                    )
                    total += 1
                    if total % 2000 == 0:
                        conn.commit()
        conn.commit()
    finally:
        wb.close()
        formula_wb.close()
    return total


def _ingest_xls(conn: sqlite3.Connection, path: Path, doc_id: int) -> int:
    """Read BIFF cells and their stored results; xlrd does not calculate formulas."""
    import xlrd
    from openpyxl.utils import get_column_letter

    book = xlrd.open_workbook(str(path), on_demand=True, formatting_info=True)
    total = 0
    try:
        for sheet in book.sheets():
            for row_index in range(sheet.nrows):
                cells: dict[str, dict[str, Any]] = {}
                for col_index in range(sheet.row_len(row_index)):
                    cell = sheet.cell(row_index, col_index)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        parsed = xlrd.xldate.xldate_as_datetime(value, book.datemode)
                        value = (parsed.date().isoformat() if parsed.time() == dt.time()
                                 else parsed.isoformat())
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = None
                    elif cell.ctype == xlrd.XL_CELL_NUMBER and value.is_integer():
                        value = int(value)
                    xf = book.xf_list[cell.xf_index]
                    number_format = book.format_map[xf.format_key].format_str
                    entry: dict[str, Any] = {"value": value, "number_format": number_format}
                    if cell.ctype == xlrd.XL_CELL_ERROR:
                        entry["warning"] = f"Excel error cell: {xlrd.error_text_from_code.get(cell.value, cell.value)}"
                    cells[get_column_letter(col_index + 1)] = entry
                if cells:
                    conn.execute(
                        "INSERT INTO sheet_rows(document_id,sheet,row_num,cells_json) VALUES(?,?,?,?)",
                        (doc_id, sheet.name, row_index + 1, json.dumps(cells, ensure_ascii=False)),
                    )
                    total += 1
                    if total % 2000 == 0:
                        conn.commit()
        conn.commit()
    finally:
        book.release_resources()
    return total


def _ingest_docx(conn: sqlite3.Connection, path: Path, doc_id: int) -> int:
    """Index body paragraphs and table rows in their original document order."""
    from docx import Document
    from docx.table import Table

    document = Document(path)
    paragraph_number = table_number = ordinal = total = 0
    section = ""
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            table_number += 1
            for row_number, row in enumerate(block.rows, 1):
                ordinal += 1
                text = " | ".join(cell.text.strip() for cell in row.cells)
                if not text.strip(" |"):
                    continue
                location = f"table {table_number} | row {row_number}"
                conn.execute(
                    "INSERT INTO word_blocks(document_id,ordinal,kind,location,section,text) "
                    "VALUES(?,?,?,?,?,?)", (doc_id, ordinal, "table", location, section, text),
                )
                total += 1
        else:
            paragraph_number += 1
            ordinal += 1
            text = block.text.strip()
            if not text:
                continue
            if block.style and block.style.name.lower().startswith("heading"):
                section = text
            conn.execute(
                "INSERT INTO word_blocks(document_id,ordinal,kind,location,section,text) "
                "VALUES(?,?,?,?,?,?)",
                (doc_id, ordinal, "paragraph", f"paragraph {paragraph_number}", section, text),
            )
            total += 1
        if total and total % 2000 == 0:
            conn.commit()
    conn.commit()
    return total


def ingest_folder(conn: sqlite3.Connection, folder: str | Path) -> dict[str, Any]:
    directory = Path(folder).expanduser().resolve()
    if not directory.is_dir():
        return {"indexed": 0, "failed": 0, "files": [], "error": "Not a readable directory"}
    # Maintain exactly one active document corpus per DB: reviewers may switch folders.
    active = conn.execute("SELECT value FROM metadata WHERE key='active_folder'").fetchone()
    if active is None or active[0] != str(directory):
        conn.execute("DELETE FROM documents")
        conn.execute("INSERT INTO metadata(key,value) VALUES('active_folder',?) "
                     "ON CONFLICT(key) DO UPDATE SET value=excluded.value", (str(directory),))
        conn.commit()
    candidates = _discover(directory)
    supported_names = {str(p.relative_to(directory)) for p in candidates
                       if p.suffix.lower() in SUPPORTED | KNOWN_UNSUPPORTED}
    # Files removed since the previous indexing run cannot remain searchable.
    for existing in conn.execute("SELECT id,path FROM documents").fetchall():
        if existing["path"] not in supported_names:
            conn.execute("DELETE FROM documents WHERE id=?", (existing["id"],))
    conn.commit()
    if not candidates:
        return {"indexed": 0, "failed": 0, "files": [], "error": "Folder is empty"}
    max_bytes = int(os.getenv("SYDEN_MAX_FILE_MB", "100")) * 1024 * 1024
    reports = []
    for path in candidates:
        suffix = path.suffix.lower()
        if suffix not in SUPPORTED | KNOWN_UNSUPPORTED:
            continue
        relative = str(path.relative_to(directory))
        stat = path.stat()
        previous = conn.execute(
            "SELECT status,file_size,mtime_ns FROM documents WHERE path=?", (relative,)
        ).fetchone()
        if (previous is not None and previous["status"] == "ready"
                and previous["file_size"] == stat.st_size
                and previous["mtime_ns"] == stat.st_mtime_ns):
            reports.append({"file": relative, "status": "skipped", "reason": "unchanged"})
            continue
        doc_id = replace_document(conn, relative, suffix.removeprefix("."))
        try:
            warnings: list[str] = []
            if suffix in KNOWN_UNSUPPORTED:
                raise ValueError(f"Unsupported Excel type {suffix}; convert to .xlsx")
            if path.stat().st_size > max_bytes:
                raise ValueError(f"File exceeds configured size limit ({max_bytes} bytes)")
            if suffix == ".pdf":
                units, warnings = _ingest_pdf(conn, path, doc_id)
            elif suffix == ".xls":
                units = _ingest_xls(conn, path, doc_id)
            elif suffix == ".docx":
                units = _ingest_docx(conn, path, doc_id)
            else:
                units = _ingest_workbook(conn, path, doc_id)
            if units == 0:
                raise ValueError("File contained no extractable text or spreadsheet rows")
            if path.stat().st_size != stat.st_size or path.stat().st_mtime_ns != stat.st_mtime_ns:
                raise ValueError("File changed while it was being indexed; retry ingestion")
            conn.execute("UPDATE documents SET file_size=?,mtime_ns=? WHERE id=?",
                         (stat.st_size, stat.st_mtime_ns, doc_id))
            mark_document(conn, doc_id, "ready", "; ".join(warnings))
            reports.append({"file": relative, "status": "ready", "units": units,
                            "warnings": warnings})
        except Exception as exc:  # noqa: BLE001 - isolate each corrupt document
            # Remove partially extracted data and keep a transparent failure record.
            conn.execute("DELETE FROM passages WHERE document_id=?", (doc_id,))
            conn.execute("DELETE FROM sheet_rows WHERE document_id=?", (doc_id,))
            conn.execute("DELETE FROM pdf_boq_rows WHERE document_id=?", (doc_id,))
            conn.execute("DELETE FROM word_blocks WHERE document_id=?", (doc_id,))
            mark_document(conn, doc_id, "failed", f"{type(exc).__name__}: {exc}")
            reports.append({"file": relative, "status": "failed", "error": str(exc)})
    return {
        "indexed": sum(r["status"] == "ready" for r in reports),
        "skipped": sum(r["status"] == "skipped" for r in reports),
        "failed": sum(r["status"] == "failed" for r in reports),
        "files": reports,
        "error": "No supported documents found" if not reports else None,
    }
