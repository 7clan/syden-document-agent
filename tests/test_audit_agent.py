"""Numerical answers must come from deterministic results, not plausible model prose."""
from __future__ import annotations

import sys
from types import SimpleNamespace

from openpyxl import Workbook

from syden_agent.agent import (
    _ask_openai,
    _final_answer,
    _needs_numeric_query,
    _question_filter,
    _supported_prose,
)
from syden_agent.db import connect
from syden_agent.ingest import ingest_folder


def _indexed_boq(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Electrical"
    sheet.append(["Item", "Total Amount", "Currency"])
    sheet.append(["Transformer", 1_500_000, "SAR"])
    sheet.append(["Cable", 1_000_000, "SAR"])
    book.save(tmp_path / "BOQ.xlsx")
    conn = connect(tmp_path / "index.db")
    assert ingest_folder(conn, tmp_path)["indexed"] == 1
    return conn


def test_numeric_intent_paraphrases_are_guarded():
    for question in (
        "List all cost items above SAR 1 million.",
        "Which equipment exceeds 1,000,000 SAR?",
        "Find electrical equipment costing more than SAR 1 million.",
        "Find equipment with prices greater than 1000000.",
        "Which BOQ items have the highest costs?",
    ):
        assert _needs_numeric_query(question), question
    assert _question_filter("In 2026, which equipment exceeds SAR 1 million?", {})["value"] == 1_000_000


def test_unrelated_real_citation_cannot_support_fabricated_budget():
    conn = connect(":memory:")
    conn.execute("INSERT INTO documents(path,kind,status) VALUES('contract.pdf','pdf','ready')")
    conn.execute("INSERT INTO passages(document_id,page,chunk,text) "
                 "VALUES(1,1,0,'Scope: install one generator')")
    answer = _final_answer(conn, "The project budget is SAR 999 million.",
                           ["contract.pdf | page 1 | chunk 0"],
                           [{"name": "search_documents", "ok": True}])
    assert "999" not in answer.text
    assert "could not be verified" in answer.text
    reversed_obligation = _final_answer(conn, "The generator is excluded.",
                                        ["contract.pdf | page 1 | chunk 0"],
                                        [{"name": "search_documents", "ok": True}])
    assert "could not be verified" in reversed_obligation.text


def test_source_location_numbers_are_not_treated_as_claimed_amounts():
    conn = connect(":memory:")
    conn.execute("INSERT INTO documents(path,kind,status) VALUES('smoke.pdf','pdf','ready')")
    conn.execute("INSERT INTO passages(document_id,page,chunk,text) "
                 "VALUES(1,1,0,'The synthetic project code is ORCHID-7429.')")
    citation = "smoke.pdf | page 1 | chunk 0"
    assert _supported_prose(conn, f"Project code: ORCHID-7429. {citation}", [citation])
    assert _supported_prose(conn, f"Project code: ORCHID-7429【search_documents 1†citation:{citation}】",
                            [citation])
    answer = _final_answer(conn, "Project code: ORCHID-7429【search_documents 1】",
                           [citation], [{"name": "search_documents", "ok": True}])
    assert answer.sources == [citation]
    assert "【" not in answer.text


def test_failed_document_marks_numeric_and_date_answers_incomplete(tmp_path):
    conn = _indexed_boq(tmp_path)
    conn.execute("INSERT INTO documents(path,kind,status,error) "
                 "VALUES('unreadable.pdf','pdf','failed','encrypted')")
    conn.commit()
    numeric = _final_answer(conn, "", [], [{"name": "query_spreadsheet", "ok": True}],
                            "List all cost items above SAR 1 million.",
                            {"operator": "gt", "value": 1_000_000, "currency": "SAR"})
    assert numeric.structured["incomplete"] is True
    assert "unreadable.pdf" in numeric.text
    dates = _final_answer(conn, "", [], [{"name": "extract_dates", "ok": True}],
                          "Extract all dates and deliverables.", date_called=True)
    assert "document coverage is incomplete" in dates.text
    assert "unreadable.pdf" in dates.text


def test_numeric_answer_ignores_fabricated_model_amount(tmp_path):
    conn = _indexed_boq(tmp_path)
    answer = _final_answer(
        conn, "Transformer costs SAR 999 million.", ["BOQ.xlsx | document"],
        [{"name": "query_spreadsheet", "ok": True}],
        "Which equipment exceeds 1,000,000 SAR?",
        {"value": 2_000_000, "currency": "AED", "operator": "gt"},
    )
    assert answer.structured["total_matches"] == 1
    assert answer.structured["verified_args"]["currency"] == "SAR"
    assert answer.structured["verified_args"]["value"] == 1_000_000
    assert answer.sources == ["BOQ.xlsx | Electrical!B2"]
    assert "1500000" in answer.text
    assert "999" not in answer.text


def test_all_boq_question_overrides_model_selected_single_workbook(tmp_path):
    first = Workbook()
    first.active.title = "Electrical"
    first.active.append(["Item", "Total Amount (SAR)"])
    first.active.append(["Transformer", 1_500_000])
    first.save(tmp_path / "first.xlsx")
    second = Workbook()
    second.active.title = "Mechanical"
    second.active.append(["Item", "Line Amount (SAR)"])
    second.active.append(["Generator", 2_250_000])
    second.save(tmp_path / "second.xlsx")
    conn = connect(tmp_path / "index.db")
    assert ingest_folder(conn, tmp_path)["indexed"] == 2
    answer = _final_answer(
        conn, "Only the transformer matched.", ["first.xlsx | Electrical!B2"],
        [{"name": "query_spreadsheet", "ok": True}],
        "List all cost items above SAR 1 million in these BOQs.",
        {"file": "first.xlsx", "sheet": "Electrical", "column_hint": "Total Amount (SAR)",
         "value": 1_000_000, "currency": "SAR"},
    )
    assert answer.structured["total_matches"] == 2
    assert answer.structured["verified_args"].get("file") is None
    assert answer.structured["verified_args"].get("column_hint") is None
    assert answer.sources == ["first.xlsx | Electrical!B2", "second.xlsx | Mechanical!B2"]
    electrical = _final_answer(
        conn, "Both items matched.", [], [{"name": "query_spreadsheet", "ok": True}],
        "Find electrical equipment costing more than SAR 1 million.",
        {"value": 1_000_000, "currency": "SAR"},
    )
    assert electrical.structured["total_matches"] == 1
    assert electrical.sources == ["first.xlsx | Electrical!B2"]


def test_mocked_provider_tool_loop_uses_verified_numbers(tmp_path, monkeypatch):
    conn = _indexed_boq(tmp_path)
    tool_message = SimpleNamespace(
        tool_calls=[SimpleNamespace(id="call_1", function=SimpleNamespace(
            name="query_spreadsheet", arguments='{"value": 1000000, "currency": "SAR"}'))],
        content=None,
        model_dump=lambda exclude_none: {"role": "assistant", "tool_calls": []},
    )
    final_message = SimpleNamespace(tool_calls=None, content="Budget is SAR 999 million.")
    responses = [SimpleNamespace(choices=[SimpleNamespace(message=tool_message)]),
                 SimpleNamespace(choices=[SimpleNamespace(message=final_message)])]
    fake_client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(
        create=lambda **kwargs: responses.pop(0))))
    monkeypatch.setitem(sys.modules, "openai", SimpleNamespace(OpenAI=lambda **kwargs: fake_client))
    monkeypatch.setenv("OPENAI_API_KEY", "synthetic-test-key")
    answer = _ask_openai(conn, "List all cost items above SAR 1 million.", 3)
    assert answer.tool_calls[0]["name"] == "query_spreadsheet"
    assert answer.structured["total_matches"] == 1
    assert "999" not in answer.text


def _contract_evidence():
    conn = connect(":memory:")
    conn.execute("INSERT INTO documents(path,kind,status) "
                 "VALUES('Contract.pdf','pdf','ready')")
    passage = ("Scope of Work Contractor shall install two transformers and commission the plant. "
               "Deliverable: Submit design drawings by 15 November 2026. "
               "Exclusion: Civil excavation is excluded from this contract.")
    conn.execute("INSERT INTO passages(document_id,page,chunk,text) VALUES(1,1,0,?)",
                 (passage,))
    return conn, "Contract.pdf | page 1 | chunk 0"


def test_scope_verifier_accepts_supported_paraphrases():
    conn, cite = _contract_evidence()
    valid = (
        ("The project scope includes the installation of two transformers and "
         "commissioning of the plant. Civil excavation is not included."),
        ("According to the contract, the work comprises installation of two "
         "transformers and plant commissioning, while civil excavation is excluded."),
        ("The contractor will install two transformers and commission the plant. "
         "Civil excavation is excluded."),
        ("The contractor will install two transformers and commission the plant "
         "without civil excavation."),
        "The contractor must submit design drawings by 15 November 2026.",
    )
    for claim in valid:
        assert _supported_prose(conn, claim, [cite]), claim
        answer = _final_answer(conn, claim, [cite],
                               [{"name": "search_documents", "ok": True}],
                               "Summarize the scope of work, including exclusions.")
        assert answer.text == claim
        assert answer.sources == [cite]


def test_scope_verifier_rejects_unsupported_claims_and_reversed_obligations():
    conn, cite = _contract_evidence()
    invalid = (
        "The contractor must perform civil excavation.",
        "The contractor must not install two transformers.",
        "The contractor shall not commission the plant.",
        "Civil excavation is included in the scope of work.",
        "The contractor must install three transformers.",
        "The contractor must install two transformers and three generators.",
        "The contract excludes commissioning the plant.",
        "The budget is SAR 50 million.",
        "The contractor will install two transformers and deliver medical equipment.",
    )
    for claim in invalid:
        assert not _supported_prose(conn, claim, [cite]), claim


def test_unsupported_summary_falls_back_to_verified_excerpts():
    conn, cite = _contract_evidence()
    claim = "The contractor must perform civil excavation and install two transformers."
    answer = _final_answer(conn, claim, [cite],
                           [{"name": "search_documents", "ok": True}],
                           "Summarize the scope of work, including exclusions.")
    assert "could not be verified" in answer.text
    assert "Verified source excerpts" in answer.text
    assert "install two transformers" in answer.text
    assert "Civil excavation is excluded" in answer.text
    assert claim not in answer.text
    assert answer.sources == [cite]


def test_prose_verification_uses_clause_level_evidence_not_cross_clause_overlap():
    conn, cite = _contract_evidence()
    assert not _supported_prose(
        conn, "The contractor shall commission civil excavation.", [cite])
    assert not _supported_prose(
        conn, "The contractor shall submit plant drawings by 1 November 2026.", [cite])


def test_scope_verifier_accepts_sectioned_summary_and_checks_exclusion_heading():
    conn, cite = _contract_evidence()
    assert _supported_prose(conn,
                            "Scope of Work:\n- Installation of two transformers and "
                            "commissioning of the plant.\nExclusions:\n- Civil excavation.",
                            [cite])
    assert not _supported_prose(conn,
                                "Scope of Work:\n- Civil excavation.", [cite])


def test_contract_fallback_sources_exclude_unrelated_document_list_entries():
    conn, cite = _contract_evidence()
    conn.execute("INSERT INTO documents(path,kind,status) "
                 "VALUES('BOQ.xlsx','xlsx','ready')")
    docs = ["BOQ.xlsx | document", "Contract.pdf | document", cite]
    answer = _final_answer(
        conn, "The contractor must perform civil excavation.", docs,
        [{"name": "list_documents", "ok": True},
         {"name": "search_documents", "ok": True}],
        "Summarize the scope of work, including exclusions.",
    )
    assert "Verified source excerpts" in answer.text
    assert answer.sources == [cite]



def test_contract_fallback_reads_as_grounded_summary_not_raw_error():
    conn, cite = _contract_evidence()
    # Different providers may add unsupported introductory language or synonyms.
    answer = _final_answer(
        conn,
        "The contractor will undertake electrical infrastructure modernization. "
        "Civil excavation is excluded.",
        [cite], [{"name": "search_documents", "ok": True}],
        "Summarize the scope of work, including exclusions.",
    )
    assert answer.text.startswith("**Verified source excerpts")
    assert "**Scope of work:**" in answer.text
    assert "**Exclusions:**" in answer.text
    assert "install two transformers and commission the plant" in answer.text
    assert "Civil excavation is excluded" in answer.text
    assert "electrical infrastructure modernization" not in answer.text
    assert "Submit design drawings" not in answer.text  # not asked for deliverables
    assert answer.sources == [cite]
    assert answer.tool_calls[-1]["name"] == "answer_verification"
    assert answer.tool_calls[-1]["ok"] is False
    assert "electrical infrastructure modernization" in (
        answer.tool_calls[-1]["generated_text_preview"]
    )
    conn.close()


def test_fallback_includes_deliverables_when_explicitly_requested():
    conn, cite = _contract_evidence()
    answer = _final_answer(
        conn, "The contractor must install generators.",
        [cite], [{"name": "search_documents", "ok": True}],
        "Summarize the scope of work, exclusions and deliverables.",
    )
    assert "**Scope of work:**" in answer.text
    assert "**Exclusions:**" in answer.text
    assert "**Deliverables:**" in answer.text
    assert "15 November 2026" in answer.text
    assert "install generators" not in answer.text
    conn.close()

    
    
def test_live_gemini_scope_preview_is_verified_as_separate_claims():
    """Regression from real tool-call audit: valid cross-clause summary was rejected."""
    conn, cite = _contract_evidence()
    generated = (
        "Based on the indexed project documents (`Contract.pdf | page 1 | chunk 0`), "
        "the scope of work and exclusions are as follows:\n\n"
        "* **Scope of Work:** The Contractor shall install two transformers and "
        "commission the plant, and submit design drawings by 15 November 2026.\n"
        "* **Exclusions:** Civil excavation is excluded from this contract."
    )
    assert _supported_prose(conn, generated, [cite])
    answer = _final_answer(
        conn, generated, [cite],
        [{"name": "list_documents", "ok": True},
         {"name": "search_documents", "ok": True},
         {"name": "get_evidence", "ok": True}],
        "Summarize the scope of work, including exclusions.",
    )
    assert answer.text == generated
    assert answer.sources == [cite]
    assert not any(call["name"] == "answer_verification" for call in answer.tool_calls)
    conn.close()


def test_coordinated_summary_does_not_launder_unsupported_obligations():
    conn, cite = _contract_evidence()
    intro = ("Based on the indexed project documents (`Contract.pdf | page 1 | chunk 0`), "
             "the scope of work and exclusions are as follows:\n")
    invalid = (
        ("* **Scope of Work:** The Contractor shall install three transformers and "
         "submit design drawings by 15 November 2026."),
        ("* **Scope of Work:** The Contractor shall install two transformers "
         "and perform civil excavation."),
        ("* **Scope of Work:** The Contractor shall install two transformers "
         "and submit design drawings by 16 November 2026."),
        ("* **Scope of Work:** The Contractor shall install two transformers "
         "and commission civil excavation."),
        "* **Exclusions:** Commissioning the plant is excluded.",
    )
    for claim in invalid:
        assert not _supported_prose(conn, intro + claim, [cite]), claim
    conn.close()


def test_real_gemini_variation_passes_without_loosening_fact_checks():
    conn, cite = _contract_evidence()
    response = (
        "Based on the indexed documents, here is the summary of the scope of work "
        "and its exclusions:\n\n"
        "* **Scope of Work:** The contractor shall install two transformers and "
        "commission the plant, along with submitting design drawings by 15 November 2026. "
        "(*Source: Contract.pdf | page 1 | chunk 0*)\n"
        "* **Exclusions:** Civil excavation is excluded from this contract. "
        "(*Source: Contract.pdf | page 1 | chunk 0*)"
    )
    assert _supported_prose(conn, response, [cite])
    result = _final_answer(conn, response, [cite], [{"name": "get_evidence", "ok": True}],
                           "Summarize the scope of work, including exclusions.")
    assert result.text == response
    assert result.sources == [cite]
    conn.close()


def test_real_groq_variation_accepts_headings_unicode_spaces_and_source_footer():
    conn, cite = _contract_evidence()
    response = (
        "**Scope of Work (Contract.pdf – page 1)**  \n"
        "- Contractor is responsible for **installing two transformers** and "
        "**commissioning the plant**.  \n"
        "- Deliverable: **Submit design drawings** by **15 November 2026**.  \n\n"
        "**Exclusions**  \n"
        "- **Civil excavation** is explicitly excluded from the contract.  \n\n"
        "*Source: Contract.pdf, page 1, chunk 0.*"
    )
    assert _supported_prose(conn, response, [cite])
    result = _final_answer(conn, response, [cite], [{"name": "get_evidence", "ok": True}],
                           "Summarize the scope of work, including exclusions.")
    assert result.text == response
    assert result.sources == [cite]
    conn.close()


def test_citation_decoration_cannot_hide_invented_work_or_wrong_dates():
    conn, cite = _contract_evidence()
    intro = ("Based on the indexed documents, here is the summary of the scope "
             "of work and its exclusions:\n")
    for wrong in (
        "- The contractor shall install four transformers and commission the plant.",
        "- The contractor shall install two transformers, along with building a hospital.",
        ("- The contractor shall install two transformers, along with submitting "
         "design drawings by 16 November 2026."),
        "- Civil excavation is included as required work.",
        "- Commissioning the plant is excluded.",
    ):
        candidate = intro + wrong + "\n*Source: Contract.pdf, page 1, chunk 0.*"
        assert not _supported_prose(conn, candidate, [cite]), wrong
    conn.close()


def test_explicit_contract_citation_omits_unrelated_search_hit():
    conn, contract_cite = _contract_evidence()
    conn.execute("INSERT INTO documents(path,kind,status) "
                 "VALUES('BOQ.xlsx','xlsx','ready')")
    conn.execute("INSERT INTO sheet_rows(document_id,sheet,row_num,cells_json) "
                 "VALUES(2,'Electrical BOQ',2,?)", (
                     '{"A":{"value":"EL-01"},"B":{"value":"Transformers"},'
                     '"C":{"value":1500000}}',))
    boq_cite = "BOQ.xlsx | Electrical BOQ | row 2"
    response = (
        "The scope of work includes the installation of two transformers and "
        "commissioning of the plant. The contractor is also required to submit "
        "design drawings by 15 November 2026.\n\n"
        "Civil excavation is excluded from this contract.\n\n"
        "Source: Contract.pdf | page 1 | chunk 0"
    )
    answer = _final_answer(
        conn, response, [contract_cite, boq_cite],
        [{"name": "search_documents", "ok": True}],
        "Summarize the scope of work, including exclusions.",
    )
    assert answer.text == response
    assert answer.sources == [contract_cite]
    conn.close()


def test_explicit_citation_does_not_discard_other_required_evidence():
    conn, contract_cite = _contract_evidence()
    conn.execute("INSERT INTO documents(path,kind,status) "
                 "VALUES('Annex.pdf','pdf','ready')")
    conn.execute("INSERT INTO passages(document_id,page,chunk,text) "
                 "VALUES(2,1,0,'Contractor shall inspect a generator.')")
    annex_cite = "Annex.pdf | page 1 | chunk 0"
    response = (
        "The contractor shall install two transformers. "
        "The contractor shall inspect a generator.\n"
        "Source: Contract.pdf | page 1 | chunk 0"
    )
    answer = _final_answer(
        conn, response, [contract_cite, annex_cite],
        [{"name": "search_documents", "ok": True}],
        "Summarize the obligations.",
    )
    assert answer.text == response
    # A single explicitly cited source cannot support the second obligation.
    assert answer.sources == [contract_cite, annex_cite]
    conn.close()
