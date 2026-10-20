"""Held-out, generated documents expose both verifier errors and its semantic limit."""
from __future__ import annotations

import io
import json

from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from syden_agent.claim_verifier import validate_claims
from syden_agent.db import connect
from syden_agent.ingest import ingest_folder
from syden_agent.tools import get_evidence, query_spreadsheet, search_documents


def _corpus(tmp_path):
    folder = tmp_path / "unseen"
    folder.mkdir()

    pdf = io.BytesIO()
    page = canvas.Canvas(pdf)
    page.drawString(50, 750, "Variation VO-19 sets a revised ceiling of SAR 2,250,000.")
    page.drawString(50, 730, "The architect handed the drawings to the subcontractor.")
    page.save()
    (folder / "Variation.pdf").write_bytes(pdf.getvalue())

    minutes = Document()
    minutes.add_heading("Delivery decisions", level=1)
    minutes.add_paragraph("VO-19 requires fire panel delivery by 12 June 2027.")
    table = minutes.add_table(rows=0, cols=3)
    for values in (("Gate", "Decision", "Due"),
                   ("TG-41", "Approved", "12 June 2027"),
                   ("TG-42", "Not approved", "18 June 2027")):
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    minutes.save(folder / "Minutes.docx")

    book = Workbook()
    sheet = book.active
    sheet.title = "Gate Register"
    sheet["B4"] = "Approval"
    sheet["C4"] = "Activity"
    sheet["D4"] = "Gate"
    sheet["E4"] = "Due"
    sheet["F4"] = "Total Amount (SAR)"
    for row, values in {
        5: ("Approved", "Fire panel delivery", "TG-41", "2027-06-12", 1_250_000),
        6: ("Not approved", "Pump trial", "TG-42", "2027-06-18", 950_000),
        8: ("Approval", "Activity", "Gate", "Due", "Total Amount (SAR)"),
        9: ("Pending", "Retest", "TG-43", "2027-07-01", 1_100_000),
    }.items():
        for column, value in zip("BCDEF", values):
            sheet[f"{column}{row}"] = value
    book.save(folder / "Gates.xlsx")

    conn = connect(tmp_path / "index.db")
    assert ingest_folder(conn, folder)["indexed"] == 3
    return conn


def _verify(conn, claim, supports, *, observed=None, question=""):
    payload = {"answer": claim, "insufficient": False, "claims": [{
        "text": claim,
        "supports": [{"citation": citation, "quote": quote}
                     for citation, quote in supports],
    }]}
    return validate_claims(conn, json.dumps(payload),
                           observed if observed is not None else
                           [citation for citation, _ in supports], question)


def test_generated_sources_have_real_pdf_word_and_excel_locations(tmp_path):
    conn = _corpus(tmp_path)
    pdf = "Variation.pdf | page 1 | chunk 0"
    word = "Minutes.docx | paragraph 2"
    table = "Minutes.docx | table 1 | row 2"
    row = "Gates.xlsx | Gate Register | row 5"
    assert "VO-19" in get_evidence(conn, pdf)["text"]
    assert "12 June 2027" in get_evidence(conn, word)["text"]
    assert "TG-41" in get_evidence(conn, table)["text"]
    assert get_evidence(conn, row)["row"]["D"] == "TG-41"
    assert any(hit["citation"] == pdf for hit in search_documents(
        conn, "revised ceiling")["results"])


def test_held_out_checkable_claims_count_false_rejections_and_acceptances(tmp_path):
    conn = _corpus(tmp_path)
    pdf = "Variation.pdf | page 1 | chunk 0"
    table = "Minutes.docx | table 1 | row 2"
    first = "Gates.xlsx | Gate Register | row 5"
    second = "Gates.xlsx | Gate Register | row 6"
    quote = "Variation VO-19 sets a revised ceiling of SAR 2,250,000."
    cases = [
        (True, "VO-19 has a revised ceiling of SAR 2,250,000.", [(pdf, quote)]),
        (True, "TG-41 is approved for fire panel delivery on 12 June 2027.",
         [(first, "")]),
        (True, "TG-41 was approved on 12 June 2027.",
         [(table, "TG-41 | Approved | 12 June 2027")]),
        (False, "VO-19 has a revised ceiling of SAR 2,500,000.", [(pdf, quote)]),
        (False, "TG-41 is due 18 June 2027.", [(first, "")]),
        (False, "TG-42 was approved.", [(second, "")]),
        (False, "TG-41 is approved.", [(second, "")]),
        (False, "VO-19 has a revised ceiling of SAR 2,250,000.",
         [(pdf, "VO-19 ceiling is SAR 2,250,000.")]),
    ]
    false_rejections = false_acceptances = 0
    errors = []
    for true_claim, claim, supports in cases:
        result = _verify(conn, claim, supports)
        accepted = result["claims"][0]["status"] != "unverified"
        false_rejections += bool(true_claim and not accepted)
        false_acceptances += bool(not true_claim and accepted)
        if accepted != true_claim:
            errors.append((claim, result["claims"][0]["status"],
                           result["claims"][0]["reason"]))
    assert false_rejections == 0, errors
    assert false_acceptances == 0, errors


def test_two_named_documents_keep_separate_claim_scopes(tmp_path):
    conn = _corpus(tmp_path)
    pdf = "Variation.pdf | page 1 | chunk 0"
    word = "Minutes.docx | paragraph 2"
    question = ("What ceiling does Variation.pdf state, and what delivery date "
                "does Minutes.docx state?")
    first = _verify(conn, "The ceiling is SAR 2,250,000.",
                    [(pdf, "Variation VO-19 sets a revised ceiling of SAR 2,250,000.")],
                    question=question)
    second = _verify(conn, "Fire panel delivery is due 12 June 2027.",
                     [(word, "VO-19 requires fire panel delivery by 12 June 2027.")],
                     question=question)
    wrong = _verify(conn, "Fire panel delivery is due 12 June 2027.",
                    [(pdf, "Variation VO-19 sets a revised ceiling of SAR 2,250,000.")],
                    question=question)
    assert first["status"] == "evidence_linked"
    assert second["status"] == "evidence_linked"
    assert wrong["status"] == "unverified"


def test_unretrieved_source_and_fabricated_word_quote_are_rejected(tmp_path):
    conn = _corpus(tmp_path)
    word = "Minutes.docx | paragraph 2"
    claim = "VO-19 requires fire panel delivery by 12 June 2027."
    real_quote = "VO-19 requires fire panel delivery by 12 June 2027."
    unobserved = _verify(conn, claim, [(word, real_quote)], observed=[])
    invented = _verify(conn, claim, [(word, "VO-19 was approved by the client.")])
    assert unobserved["status"] == "unverified"
    assert invented["status"] == "unverified"


def test_financial_query_uses_original_cells_not_model_amounts(tmp_path):
    conn = _corpus(tmp_path)
    result = query_spreadsheet(conn, file="Gates.xlsx", currency="SAR",
                               operator="gt", value=1_000_000)
    assert result["total_matches"] == 2, result
    assert result["matched_amount_sum"] == "2350000"
    assert {item["citation"] for item in result["matches"]} == {
        "Gates.xlsx | Gate Register!F5", "Gates.xlsx | Gate Register!F9"}


def test_authentic_quote_cannot_prove_role_direction(tmp_path):
    """Known semantic false acceptances must remain labeled as interpretation."""
    conn = _corpus(tmp_path)
    citation = "Variation.pdf | page 1 | chunk 0"
    wrong = _verify(
        conn, "The subcontractor handed the drawings to the architect.",
        [(citation, "The architect handed the drawings to the subcontractor.")],
    )
    assert wrong["status"] == "evidence_linked"
    assert wrong["claims"][0]["reason"] == "Interpretation is not independently proven"
    wrong_row = _verify(conn, "TG-41 is a pump trial.",
                        [("Gates.xlsx | Gate Register | row 5", "")])
    assert wrong_row["status"] == "evidence_linked"
    assert wrong_row["claims"][0]["reason"] == "Interpretation is not independently proven"
