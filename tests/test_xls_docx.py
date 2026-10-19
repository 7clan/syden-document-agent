"""Synthetic legacy Excel and Word documents; no provider credentials required."""
from __future__ import annotations

import datetime as dt
import struct

import pytest
import xlrd
import xlwt
from docx import Document
from openpyxl import Workbook

from syden_agent.db import connect
from syden_agent.ingest import ingest_folder
from syden_agent.tools import (
    extract_dates,
    get_evidence,
    inspect_document,
    query_spreadsheet,
    search_documents,
)


def _workbooks(folder):
    legacy = xlwt.Workbook()
    modern = Workbook()
    for sheet_name, amounts in (("Electrical SAR", (1_500_000, 1_000_000, 2_250_000)),
                                ("Mechanical AED", (2_000_000, 3_000_000))):
        old = legacy.add_sheet(sheet_name)
        new = modern.active if sheet_name == "Electrical SAR" else modern.create_sheet(sheet_name)
        new.title = sheet_name
        currency = "SAR" if "SAR" in sheet_name else "AED"
        rows = [["Item", f"Total Amount ({currency})"]]
        rows.extend([[f"Item {number}", amount] for number, amount in enumerate(amounts, 1)])
        if currency == "SAR":
            rows.extend([["Item", "Unit Price"], ["Item", "Line Amount (SAR)"],
                         ["Repeated", 1_250_000]])
        for row_num, values in enumerate(rows):
            for col_num, value in enumerate(values):
                old.write(row_num, col_num, value)
            new.append(values)
    legacy.save(str(folder / "costs.xls"))
    modern.save(folder / "costs.xlsx")


def test_xls_and_xlsx_numeric_pipeline_match(tmp_path):
    _workbooks(tmp_path)
    conn = connect(tmp_path / "index.db")
    assert ingest_folder(conn, tmp_path)["indexed"] == 2
    for suffix in ("xls", "xlsx"):
        sar = query_spreadsheet(conn, file=f"costs.{suffix}", currency="SAR",
                                value=1_000_000, limit=1)
        assert sar["total_matches"] == 3
        assert sar["matched_amount_sum"] == "5000000"
        assert sar["has_more"] is True
        assert sar["inspected_columns"][0]["header_row"] == 1
        second = query_spreadsheet(conn, file=f"costs.{suffix}", currency="SAR",
                                   value=1_000_000, offset=2, limit=1)
        assert second["matches"][0]["citation"] == f"costs.{suffix} | Electrical SAR!B7"
        assert get_evidence(conn, second["matches"][0]["citation"])["value"]["value"] == 1_250_000
        aed = query_spreadsheet(conn, file=f"costs.{suffix}", currency="AED", value=1_000_000)
        assert aed["total_matches"] == 2
        assert aed["matched_amount_sum"] == "5000000"
        assert {s["sheet"] for s in inspect_document(conn, f"costs.{suffix}")["sheets"]} == {
            "Electrical SAR", "Mechanical AED"}


@pytest.mark.parametrize("datemode", [0, 1])
def test_xls_date_system_and_formula_cache(tmp_path, datemode):
    legacy = xlwt.Workbook()
    legacy.dates_1904 = datemode
    sheet = legacy.add_sheet("Milestones SAR")
    sheet.write(0, 0, "Due date")
    sheet.write(0, 1, "Total Amount (SAR)")
    sheet.write(1, 0, dt.datetime(2026, 11, 15), xlwt.easyxf(num_format_str="YYYY-MM-DD"))
    sheet.write(1, 1, 1_500_000)
    sheet.write(2, 0, "Formula without cached result")
    sheet.write(2, 1, xlwt.Formula("B2+1"))
    sheet.write(3, 1, xlwt.Formula("B2+2"))
    legacy.save(str(tmp_path / "dates.xls"))
    # xlwt writes formula placeholders. Replace one BIFF FORMULA record's 8-byte cached
    # result inside this synthetic file; the second formula remains uncached.
    path = tmp_path / "dates.xls"
    data = bytearray(path.read_bytes())
    formula_record = bytes.fromhex("06001f0002000100")  # row 3, column B
    assert data.count(formula_record) == 1
    start = data.index(formula_record) + 10
    data[start:start + 8] = struct.pack("<d", 1_500_001)
    path.write_bytes(data)
    assert xlrd.open_workbook(str(tmp_path / "dates.xls")).datemode == datemode
    conn = connect(tmp_path / "index.db")
    assert ingest_folder(conn, tmp_path)["indexed"] == 1
    evidence = get_evidence(conn, "dates.xls | Milestones SAR!A2")
    assert evidence["value"]["value"].startswith("2026-11-15")
    dates = extract_dates(conn)
    assert dates["dates"][0]["normalized_date"] == "2026-11-15"
    assert dates["dates"][0]["citation"] == "dates.xls | Milestones SAR!A2"
    amounts = query_spreadsheet(conn, file="dates.xls", currency="SAR", value=1_000_000)
    assert amounts["total_matches"] == 2
    formula = get_evidence(conn, "dates.xls | Milestones SAR!B3")
    assert formula["value"]["value"] == 1_500_001
    assert get_evidence(conn, "dates.xls | Milestones SAR!B4")["value"]["value"] == ""


def test_xls_sparse_cells_keep_original_row_and_column(tmp_path):
    legacy = xlwt.Workbook()
    sheet = legacy.add_sheet("Sparse SAR")
    sheet.write(11, 1, "Description")
    sheet.write(11, 3, "Total Amount (SAR)")
    sheet.write(12, 1, "Remote item")
    sheet.write(12, 3, 1_250_000)
    legacy.save(str(tmp_path / "sparse.xls"))
    conn = connect(tmp_path / "index.db")
    assert ingest_folder(conn, tmp_path)["indexed"] == 1
    matches = query_spreadsheet(conn, file="sparse.xls", currency="SAR", value=1_000_000)
    assert matches["matches"][0]["citation"] == "sparse.xls | Sparse SAR!D13"
    assert get_evidence(conn, "sparse.xls | Sparse SAR!D13")["value"]["value"] == 1_250_000


def test_docx_order_sources_and_cross_document_search(tmp_path):
    word = Document()
    word.add_heading("Scope of Work", level=1)
    word.add_paragraph("Install transformers by 15 November 2026.")
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Deliverable"
    table.cell(0, 1).text = "Due"
    table.cell(1, 0).text = "Commission plant"
    table.cell(1, 1).text = "2026-12-01"
    word.add_heading("Exclusions", level=1)
    word.add_paragraph("Civil excavation is excluded.")
    word.save(tmp_path / "contract.docx")
    modern = Workbook()
    modern.active.append(["Item", "Total Amount (SAR)"])
    modern.active.append(["Transformer", 2_000_000])
    modern.save(tmp_path / "boq.xlsx")
    conn = connect(tmp_path / "index.db")
    assert ingest_folder(conn, tmp_path)["indexed"] == 2
    preview = inspect_document(conn, "contract.docx")["blocks_preview"]
    assert [p["location"] for p in preview] == [
        "paragraph 1", "paragraph 2", "table 1 | row 1", "table 1 | row 2",
        "paragraph 3", "paragraph 4"]
    matches = search_documents(conn, "transformer")
    assert {m["citation"] for m in matches["results"]} == {
        "contract.docx | paragraph 2", "boq.xlsx | Sheet | row 2"}
    assert get_evidence(conn, "contract.docx | table 1 | row 2")["text"] == (
        "Commission plant | 2026-12-01")
    assert get_evidence(conn, "contract.docx | paragraph 4")["section"] == "Exclusions"
    assert "error" in get_evidence(conn, "contract.docx | paragraph 99")
    assert {d["citation"] for d in extract_dates(conn)["dates"]} == {
        "contract.docx | paragraph 2", "contract.docx | table 1 | row 2"}


def test_bad_and_empty_new_formats_do_not_hide_valid_documents(tmp_path):
    (tmp_path / "bad.xls").write_bytes(b"bad xls")
    (tmp_path / "bad.docx").write_bytes(b"bad docx")
    empty = xlwt.Workbook()
    empty.add_sheet("Empty")
    empty.save(str(tmp_path / "empty.xls"))
    Document().save(tmp_path / "empty.docx")
    good = Document()
    good.add_paragraph("Valid contract text")
    good.save(tmp_path / "good.docx")
    conn = connect(tmp_path / "index.db")
    report = ingest_folder(conn, tmp_path)
    assert report["indexed"] == 1
    assert report["failed"] == 4
    assert search_documents(conn, "Valid contract")["total_matches"] == 1
    assert search_documents(conn, "bad xls")["total_matches"] == 0


def test_word_reindex_removes_stale_paragraphs(tmp_path):
    path = tmp_path / "contract.docx"
    first = Document()
    first.add_paragraph("Old obligation")
    first.save(path)
    conn = connect(tmp_path / "index.db")
    assert ingest_folder(conn, tmp_path)["indexed"] == 1
    assert search_documents(conn, "Old obligation")["total_matches"] == 1
    second = Document()
    second.add_paragraph("New obligation with more words")
    second.save(path)
    assert ingest_folder(conn, tmp_path)["indexed"] == 1
    assert search_documents(conn, "Old")["total_matches"] == 0
    assert search_documents(conn, "New obligation")["total_matches"] == 1
