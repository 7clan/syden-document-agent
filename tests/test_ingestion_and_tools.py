"""Offline tests: no API key or internet required."""
from __future__ import annotations

from openpyxl import Workbook

from syden_agent.db import connect
from syden_agent.ingest import ingest_folder
from syden_agent.tools import (
    extract_dates,
    get_evidence,
    inspect_document,
    list_documents,
    query_spreadsheet,
    search_documents,
)


def _make_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Electrical BOQ"
    ws.append(["Item", "Description", "Total Amount (SAR)", "Unit Price"])
    ws.append(["EL-01", "Transformer", 1_500_000, 1_500])
    ws.append(["EL-02", "Copper cable", 500_000, 12])
    ws.append(["EL-03", "Generator", 1_000_000, 100_000])
    ws.append(["EL-04", "Expensive panel", 2_250_000, 200])
    ws.append(["EL-05", "Formula missing cached value", "=1000000+5", 0])
    wb.save(path)


def test_index_and_threshold_are_exact(tmp_path):
    _make_workbook(tmp_path / "boq.xlsx")
    conn = connect(tmp_path / "index.db")
    report = ingest_folder(conn, tmp_path)
    assert report["indexed"] == 1
    assert report["failed"] == 0
    docs = list_documents(conn)
    assert len(docs["documents"]) == 1
    assert inspect_document(conn, "boq.xlsx")["sheets"][0]["sheet"] == "Electrical BOQ"
    result = query_spreadsheet(conn, column_hint="total amount", operator="gt",
                               value=1_000_000, currency="SAR")
    assert result["total_matches"] == 2
    assert [item["row"] for item in result["matches"]] == [2,5]
    assert result["matches"][0]["citation"] == "boq.xlsx | Electrical BOQ!C2"
    assert result["inspected_columns"][0]["currency_verified"] is True
    cite = get_evidence(conn, result["matches"][0]["citation"])
    assert cite["value"]["value"] == 1_500_000
    conn.close()


def test_pagination_and_strict_greater_than(tmp_path):
    _make_workbook(tmp_path / "boq.xlsx")
    conn = connect(tmp_path / "db.sqlite")
    ingest_folder(conn, tmp_path)
    more = query_spreadsheet(conn, value=1_000_000, operator="gte", limit=1, offset=1)
    assert more["total_matches"] == 3
    assert more["has_more"] is True
    assert len(more["matches"]) == 1
    symbolic = query_spreadsheet(conn, value=1_000_000, operator=">", currency="SAR")
    assert symbolic["total_matches"] == 2
    conn.close()


def test_empty_folder_and_corrupt_file(tmp_path):
    conn = connect(tmp_path / "db.sqlite")
    report = ingest_folder(conn, tmp_path / "missing")
    assert report["indexed"] == 0
    assert "error" in report
    (tmp_path / "bad.pdf").write_bytes(b"not a PDF")
    (tmp_path / "old.xls").write_bytes(b"legacy binary file")
    report = ingest_folder(conn, tmp_path)
    assert report["indexed"] == 0
    assert report["failed"] == 2
    assert all(x["status"] == "failed" for x in list_documents(conn)["documents"])
    conn.close()


def test_switch_folders_clears_previous_corpus_and_removed_files(tmp_path):
    folder_a = tmp_path / "a"
    folder_b = tmp_path / "b"
    folder_a.mkdir()
    folder_b.mkdir()
    _make_workbook(folder_a / "a.xlsx")
    conn = connect(tmp_path / "db.sqlite")
    assert ingest_folder(conn, folder_a)["indexed"] == 1
    assert [x["path"] for x in list_documents(conn)["documents"]] == ["a.xlsx"]
    # Folder changes must not expose files from an earlier live-demo corpus.
    assert ingest_folder(conn, folder_b)["indexed"] == 0
    assert list_documents(conn)["documents"] == []
    _make_workbook(folder_b / "b.xlsx")
    assert ingest_folder(conn, folder_b)["indexed"] == 1
    assert [x["path"] for x in list_documents(conn)["documents"]] == ["b.xlsx"]
    (folder_b / "b.xlsx").unlink()
    assert ingest_folder(conn, folder_b)["indexed"] == 0
    assert list_documents(conn)["documents"] == []
    conn.close()


def test_repeated_headers_change_amount_column_and_preserve_duplicates(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ SAR"
    ws.append(["Item", "Description", "Total Amount (SAR)", "Unit Rate"])
    ws.append(["A", "First", 1_500_000, 5_000_000])
    for _ in range(8):
        ws.append([])
    ws.append(["Item", "Description", "Unit Rate", "Line Amount (SAR)"])
    ws.append(["B", "Second", 9_000_000, 1_250_000])
    ws.append(["B", "Second", 9_000_000, 1_250_000])
    ws.append(["C", "Boundary", 9_000_000, 1_000_000])
    wb.save(tmp_path / "changed.xlsx")
    conn = connect(tmp_path / "index.db")
    assert ingest_folder(conn, tmp_path)["indexed"] == 1
    first = query_spreadsheet(conn, currency="SAR", value=1_000_000, limit=1)
    assert first["total_matches"] == 3
    assert first["matched_amount_sum"] == "4000000"
    assert first["has_more"] is True
    assert first["matches"][0]["citation"] == "changed.xlsx | BOQ SAR!C2"
    second = query_spreadsheet(conn, currency="SAR", value=1_000_000, offset=1, limit=2)
    assert [m["citation"] for m in second["matches"]] == [
        "changed.xlsx | BOQ SAR!D12", "changed.xlsx | BOQ SAR!D13"]
    assert second["has_more"] is False
    assert get_evidence(conn, second["matches"][0]["citation"])["value"]["value"] == 1_250_000


def test_currency_is_not_inferred_from_other_amount_column(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Mixed"
    ws.append(["Item", "Total Amount (AED)", "Unit Price (SAR)"])
    ws.append(["A", 2_000_000, 100])
    wb.save(tmp_path / "mixed.xlsx")
    conn = connect(tmp_path / "index.db")
    ingest_folder(conn, tmp_path)
    sar = query_spreadsheet(conn, currency="SAR", value=1_000_000)
    assert sar["total_matches"] == 0
    assert sar["incomplete"] is False
    assert sar["warnings"] == []
    aed = query_spreadsheet(conn, currency="AED", value=1_000_000)
    assert aed["total_matches"] == 1
    assert aed["matches"][0]["citation"] == "mixed.xlsx | Mixed!B2"


def test_foreign_currency_sheet_does_not_make_sar_results_incomplete(tmp_path):
    book = Workbook()
    sar = book.active
    sar.title = "SAR Items"
    sar.append(["Item", "Extended Total (SAR)"])
    sar.append(["Included", 1_250_000])
    sar.append(["Boundary", 1_000_000])
    aed = book.create_sheet("AED Options")
    aed.append(["Item", "Extended Total (AED)"])
    aed.append(["Excluded", 2_000_000])
    book.save(tmp_path / "mixed_sheets.xlsx")
    conn = connect(tmp_path / "index.db")
    assert ingest_folder(conn, tmp_path)["indexed"] == 1

    result = query_spreadsheet(conn, currency="SAR", value=1_000_000)

    assert result["total_matches"] == 1
    assert result["matched_amount_sum"] == "1250000"
    assert result["matches"][0]["citation"] == "mixed_sheets.xlsx | SAR Items!B2"
    assert result["incomplete"] is False
    assert result["warnings"] == []


def test_sparse_rows_keep_actual_cell_addresses(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ SAR"
    ws["B12"] = "Description"
    ws["D12"] = "Total Amount (SAR)"
    ws["B13"] = "Sparse item"
    ws["D13"] = 1_200_000
    wb.save(tmp_path / "sparse.xlsx")
    conn = connect(tmp_path / "index.db")
    assert ingest_folder(conn, tmp_path)["indexed"] == 1
    result = query_spreadsheet(conn, currency="SAR", value=1_000_000)
    assert result["total_matches"] == 1
    assert result["matches"][0]["citation"] == "sparse.xlsx | BOQ SAR!D13"
    evidence = get_evidence(conn, result["matches"][0]["citation"])
    assert evidence["value"]["value"] == 1_200_000


def test_sums_are_never_combined_across_currencies(tmp_path):
    wb = Workbook()
    sar = wb.active
    sar.title = "SAR"
    sar.append(["Item", "Total Amount (SAR)"])
    sar.append(["A", 2_000_000])
    aed = wb.create_sheet("AED")
    aed.append(["Item", "Total Amount (AED)"])
    aed.append(["B", 3_000_000])
    wb.save(tmp_path / "two.xlsx")
    conn = connect(tmp_path / "index.db")
    ingest_folder(conn, tmp_path)
    result = query_spreadsheet(conn, value=1_000_000)
    assert result["total_matches"] == 2
    assert result["matched_amount_sum"] is None
    assert result["matched_amount_sums"] == {"AED": "3000000", "SAR": "2000000"}


def test_repeat_ingestion_skips_unchanged_and_updates_changed_file(tmp_path):
    path = tmp_path / "boq.xlsx"
    _make_workbook(path)
    conn = connect(tmp_path / "index.db")
    assert ingest_folder(conn, tmp_path)["indexed"] == 1
    repeat = ingest_folder(conn, tmp_path)
    assert repeat["indexed"] == 0
    assert repeat["skipped"] == 1
    assert query_spreadsheet(conn, currency="SAR")["total_matches"] == 2
    wb = Workbook()
    ws = wb.active
    ws.append(["Item", "Total Amount (SAR)"])
    ws.append(["NEW", 4_000_000])
    wb.save(path)
    updated = ingest_folder(conn, tmp_path)
    assert updated["indexed"] == 1
    result = query_spreadsheet(conn, currency="SAR")
    assert result["total_matches"] == 1
    assert result["matches"][0]["citation"] == "boq.xlsx | Sheet!B2"


def test_dates_preserve_ambiguity_and_row_context(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Milestones"
    ws.append(["Deliverable", "Due date"])
    ws.append(["Submit design drawings", "2026-11-15"])
    ws.append(["Commission plant", "03/04/26"])
    ws.append(["Submit report", "17/04/2026"])
    wb.save(tmp_path / "dates.xlsx")
    conn = connect(tmp_path / "index.db")
    ingest_folder(conn, tmp_path)
    dates = extract_dates(conn, limit=2)
    assert dates["total_matches"] == 3
    assert dates["has_more"] is True
    assert dates["dates"][0]["normalized_date"] == "2026-11-15"
    assert dates["dates"][0]["context"]["A"] == "Submit design drawings"
    assert dates["dates"][1]["ambiguous"] is True
    assert dates["dates"][1]["normalized_date"] is None
    last = extract_dates(conn, limit=1, offset=2)["dates"][0]
    assert last["normalized_date"] == "2026-04-17"
    assert get_evidence(conn, last["citation"])["value"]["value"] == "17/04/2026"


def test_spreadsheet_search_citation_resolves_to_real_row(tmp_path):
    _make_workbook(tmp_path / "boq.xlsx")
    conn = connect(tmp_path / "index.db")
    ingest_folder(conn, tmp_path)
    result = search_documents(conn, "Transformer", limit=1)
    assert result["total_matches"] == 1
    citation = result["results"][0]["citation"]
    assert citation == "boq.xlsx | Electrical BOQ | row 2"
    assert get_evidence(conn, citation)["row"]["B"] == "Transformer"
    assert "error" in get_evidence(conn, "boq.xlsx | Electrical BOQ!Z2")


def test_header_after_fifty_populated_rows_is_found(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ SAR"
    for number in range(1, 60):
        ws.append([f"Project note {number}"])
    ws.append(["Item", "Total Amount (SAR)"])
    ws.append(["LATE", 1_750_000])
    wb.save(tmp_path / "late.xlsx")
    conn = connect(tmp_path / "index.db")
    ingest_folder(conn, tmp_path)
    result = query_spreadsheet(conn, currency="SAR", value=1_000_000)
    assert result["total_matches"] == 1
    assert result["matches"][0]["citation"] == "late.xlsx | BOQ SAR!B61"
    assert not result["warnings"]


def test_large_sparse_sheet_and_missing_formula_cache(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ SAR"
    ws["B4"] = "Item"
    ws["D4"] = "Total Amount (SAR)"
    ws["B20000"] = "Remote row"
    ws["D20000"] = 1_875_000
    ws["D20000"].number_format = "#,##0.00"
    ws["B20001"] = "Uncalculated formula"
    ws["D20001"] = "=1000000+5"
    wb.save(tmp_path / "large_sparse.xlsx")
    conn = connect(tmp_path / "index.db")
    assert ingest_folder(conn, tmp_path)["indexed"] == 1
    result = query_spreadsheet(conn, currency="SAR", value=1_000_000)
    assert result["total_matches"] == 1
    assert result["matches"][0]["citation"] == "large_sparse.xlsx | BOQ SAR!D20000"
    evidence = get_evidence(conn, result["matches"][0]["citation"])
    assert evidence["value"]["number_format"] == "#,##0.00"
    formula = get_evidence(conn, "large_sparse.xlsx | BOQ SAR!D20001")["value"]
    assert formula["value"] is None
    assert formula["formula"] == "=1000000+5"
    assert "no cached result" in formula["warning"]
